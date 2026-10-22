"""Extract the supplied annual bank data without treating spreadsheet prices as live."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


SYMBOL_NAMES = {
    "1120": "مصرف الراجحي",
    "1180": "البنك الأهلي السعودي",
    "1010": "بنك الرياض",
    "1060": "البنك السعودي الأول",
    "1150": "مصرف الإنماء",
    "1050": "البنك السعودي الفرنسي",
    "1080": "البنك العربي الوطني",
    "1140": "بنك البلاد",
    "1020": "بنك الجزيرة",
    "1030": "البنك السعودي للاستثمار",
}


def extract(source: Path) -> dict:
    workbook = load_workbook(source, read_only=True, data_only=True)
    banks = []
    for sheet in workbook.worksheets:
        if sheet.title == "الودائع تحت الطلب":
            continue
        symbol = str(sheet["B11"].value).strip()
        if symbol not in SYMBOL_NAMES:
            raise ValueError(f"Unknown symbol in {sheet.title}: {symbol}")
        annual = []
        quality_flags = []
        for row in range(2, 8):
            year = int(sheet.cell(row, 2).value)
            period = str(sheet.cell(row, 3).value).strip()
            basic = float(sheet.cell(row, 11).value)
            diluted = float(sheet.cell(row, 12).value)
            if period != "FY" or year != 2018 + row:
                raise ValueError(f"Unexpected annual row {sheet.title}!{row}")
            if diluted > basic and basic >= 0:
                quality_flags.append(
                    f"ربح السهم المخفف أكبر من الأساسي في {sheet.title}، الصف {row}؛ يلزم التحقق من إفصاح البنك"
                )
            annual.append(
                {
                    "year": year,
                    "basicEps": basic,
                    "dilutedEps": diluted,
                    "netIncomeMillionSar": float(sheet.cell(row, 10).value),
                    "cells": {"basicEps": f"K{row}", "dilutedEps": f"L{row}"},
                }
            )
        book = float(sheet["B16"].value)
        banks.append(
            {
                "symbol": symbol,
                "name": SYMBOL_NAMES[symbol],
                "sourceSheet": sheet.title,
                "annual": annual,
                "workbookBookValue": {
                    "value": book,
                    "asOf": None,
                    "cell": "B16",
                },
                "qualityFlags": quality_flags,
            }
        )
    if len(banks) != 10:
        raise ValueError(f"Expected 10 banks, found {len(banks)}")
    return {
        "schemaVersion": 1,
        "source": {
            "fileName": source.name,
            "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "period": "FY 2020–2025",
            "bookValueDateKnown": False,
        },
        "banks": banks,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    result = extract(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"Extracted {len(result['banks'])} banks to {args.output}")


if __name__ == "__main__":
    main()
